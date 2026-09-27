#!/usr/bin/env python3
"""Create a review-only BioSample workbook from a supplied clinical workbook and template."""
import argparse, copy, csv, json, re
from pathlib import Path
from openpyxl import load_workbook

EMPTY = {'', '-', 'NA', 'N/A', 'UNSTATED', 'NONE'}
KNOWN = {'Sample name', 'Age', 'Sex', 'NCBI taxonomy id', 'Organism', 'Biomaterial provider'}

def norm(v):
    return re.sub(r'\s+', ' ', str(v or '').strip()).casefold()

def locate_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        found = {norm(c.value): c.column for c in row if norm(c.value) in {norm(x) for x in KNOWN}}
        if 'sample name' in found and len(found) >= 2:
            return row[0].row, found
    raise ValueError('BioSample header row not found; supply an actual template with Sample name')

def main():
    p = argparse.ArgumentParser(description=__doc__)
    for arg in ('source', 'template', 'config', 'output', 'audit'):
        p.add_argument('--'+arg, required=True)
    a = p.parse_args()
    for item in (a.source, a.template, a.config):
        if not Path(item).is_file():
            p.error(f'missing input: {item}')
    cfg = json.loads(Path(a.config).read_text(encoding='utf-8'))
    if cfg.get('approved_facts') and not isinstance(cfg['approved_facts'], dict):
        p.error('approved_facts must be an object')
    wb = load_workbook(a.source, read_only=True, data_only=True)
    src = wb[cfg.get('source_sheet', wb.sheetnames[0])]
    headers = [c.value for c in next(src.iter_rows(min_row=1, max_row=1))]
    for h in ('SAMPLE', 'mRNA', 'QC_Macrogen', 'AGE', 'SEX'):
        if h not in headers:
            p.error(f'source column missing: {h}')
    rows = [dict(zip(headers, r)) for r in src.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
    dest = load_workbook(a.template)
    sheet = dest[cfg.get('template_sheet', dest.sheetnames[0])]
    hr, columns = locate_header(sheet)
    write_row = int(cfg.get('first_data_row', hr+2))
    if any(sheet.cell(r, columns['sample name']).value not in (None, '') for r in range(write_row, sheet.max_row+1)):
        p.error('template already contains sample values; use an empty copy')
    facts = {}
    for k, assertion in cfg.get('approved_facts', {}).items():
        key = norm(k)
        if key in {'sample name', 'age', 'sex'}:
            p.error(f'approved_facts cannot override a row-derived field: {k}')
        if not isinstance(assertion, dict) or not assertion.get('evidence') or not assertion.get('value'):
            p.error(f'approved_facts.{k} needs a value and evidence citation')
        facts[key] = assertion['value']
    if any(k not in columns for k in facts):
        p.error('approved_facts contains a field absent from the supplied template')
    audit = []
    seen = set()
    if write_row <= hr:
        p.error('first_data_row must follow the header')
    for offset, source in enumerate(rows, 2):
        raw_id = str(source['SAMPLE'] or '').strip()
        q = str(source['QC_Macrogen'] or '').strip().lower()
        mrna = str(source['mRNA'] or '').strip()
        if q != 'true' or norm(mrna).upper() in EMPTY or not mrna:
            audit.append([offset, raw_id, 'EXCLUDED_QC_OR_NO_MRNA', '', 'No BioSample row generated'])
            continue
        if mrna in seen:
            audit.append([offset, raw_id, 'DUPLICATE_MRNA', mrna, 'No BioSample row generated'])
            continue
        seen.add(mrna)
        values = {'sample name': mrna}
        age = source['AGE']
        if norm(age).upper() not in EMPTY and age is not None and cfg.get('age_unit'):
            values['age'] = f'{age} {cfg["age_unit"]}'
        sex = {'M': 'male', 'F': 'female'}.get(str(source['SEX'] or '').strip().upper())
        if sex:
            values['sex'] = sex
        values.update(facts)
        for key, col in columns.items():
            cell = sheet.cell(write_row, col)
            if write_row > hr+1:
                model = sheet.cell(int(cfg.get('first_data_row', hr+2)), col)
                if model.has_style:
                    cell._style = copy.copy(model._style)
                if model.number_format:
                    cell.number_format = model.number_format
            if key in values:
                cell.value = values[key]
        missing = [k for k in columns if k not in values]
        audit.append([offset, raw_id, 'DRAFT', mrna, 'Filled: '+', '.join(sorted(values))+'; Blank: '+', '.join(sorted(missing))])
        write_row += 1
    if Path(a.output).resolve() == Path(a.template).resolve():
        p.error('output must differ from template')
    dest.save(a.output)
    with open(a.audit, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f); writer.writerow(['source_excel_row','clinical_sample','status','biosample_name','detail']);writer.writerows(audit)
    print(json.dumps({'source_rows':len(rows),'draft_rows':write_row-int(cfg.get('first_data_row', hr+2)),
                      'excluded_or_duplicate':sum(x[2]!='DRAFT' for x in audit),'header_row':hr,
                      'unfilled_template_fields':sorted(set(columns)-{'sample name','age','sex'}-set(facts))},ensure_ascii=False))

if __name__ == '__main__':
    main()
